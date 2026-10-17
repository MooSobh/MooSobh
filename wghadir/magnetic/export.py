"""Excel workbook of the processed magnetic data (README, data, day summary, crossovers, IGRF).

Author : Dr. Mohamed Sobh, LIAG (Leibniz Institute for Applied Geophysics)
Project: Wadi Ghadir gravity and magnetic survey, Eastern Desert, Egypt (2026)
"""
from openpyxl import Workbook
from openpyxl.styles import Font

from ..common.excel import write_sheet


def write_workbook(path, d, days, xs, igrf_sens, meta):
    """Write the magnetic results workbook (README, Mag_Data, Day_Summary, Crossovers, IGRF_Date_Sensitivity)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [("Wadi Ghadir ground magnetic survey – processed total field, anomaly and QC", True),
             ("Prepared by Dr. Mohamed Sobh, LIAG. Produced by the magnetic processing chain (wghadir.chains.magnetic).", False), ("", False)]
    lines += [(f"{k}: {v}", False) for k, v in meta.items()]
    lines += [("", False), ("Columns of Mag_Data", True),
              ("tmi_corr_nT = diurnally corrected total field as supplied; igrf_F_nT = IGRF-14 total field; "
               "tma_nT = tmi_corr_nT - igrf_F_nT (total-field anomaly); accepted = not a spike/non-data row; "
               "flags = QC reasons; day_seq = survey day inferred from the time stamps.", False)]
    for t, b in lines:
        ws.append([t])
        if b:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
    ws.column_dimensions["A"].width = 150
    cols = ["source_row", "day_seq", "time_hhmmss", "lat", "lon", "utm_e", "utm_n", "elev_file", "reading_nT",
            "diurnal_corr_nT", "tmi_corr_nT", "igrf_F_nT", "igrf_inc_deg", "igrf_dec_deg", "tma_nT",
            "hampel_dev_nT", "accepted", "flags"]
    write_sheet(wb.create_sheet("Mag_Data"), d[cols], {"lat": "0.0000000", "lon": "0.0000000", "utm": "0.00",
                                                       "_nT": "0.00", "_deg": "0.000"})
    write_sheet(wb.create_sheet("Day_Summary"), days)
    write_sheet(wb.create_sheet("Crossovers"), xs)
    write_sheet(wb.create_sheet("IGRF_Date_Sensitivity"), igrf_sens)
    wb.save(path)
