"""Ground magnetic data of the Wadi Ghadir survey: QC, IGRF removal, maps and RTP.

Input: one workbook with columns lat, long, x, y, E, Reading, Time, Corrected data
(Total Intensity); "Corrected data" is the diurnally corrected total field supplied
by the field team (Corrected - Reading = diurnal correction, 0 to -170 nT).

Chain
  1  load every row, keep the source row number, flag non-data rows
  2  day segments from the time stamps (the file carries no dates)
  3  despike: Hampel filter within each day (window 11 readings ~ 28 m)
  4  IGRF-14 total field at each reading (ppigrf) -> total-field anomaly
  5  crossover check between days (readings < 5 m apart on different days)
  6  gridding as for gravity: UTM 36N, block median, damped biharmonic spline,
     blanked > MASK_KM from data
  7  reduction to the pole (Harmonica, FFT) assuming induced magnetisation along
     the IGRF field direction; tapered and padded grid to limit edge effects
"""
import datetime as dt
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.patheffects as pe  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import textwrap  # noqa: E402

from matplotlib.colors import BoundaryNorm  # noqa: E402

from . import maps  # noqa: E402

SURVEY_DATE = dt.datetime(2026, 1, 20)   # ASSUMED: the file has times but no dates
HAMPEL_WINDOW = 11
HAMPEL_K = 6.0
HAMPEL_FLOOR_NT = 30.0
CROSSOVER_M = 5.0
GRID_SPACING_M = 50.0
BLOCK_M = 100.0
MASK_KM = 2.0
TAPER_KM = 3.0


# ----------------------------------------------------------------------------- load and QC
def load(path):
    import pyproj

    raw = pd.read_excel(path, sheet_name=0, header=0)
    raw = raw.iloc[:, :8]
    raw.columns = ["lat", "lon", "x_file", "y_file", "elev_file", "reading_nT", "time_hhmmss", "tmi_corr_nT"]
    raw.insert(0, "source_row", np.arange(2, len(raw) + 2))  # Excel row number
    d = raw.copy()
    for c in d.columns[1:]:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    d["flags"] = ""
    nd = d["time_hhmmss"].isna() | d["lon"].isna() | d["lat"].isna()
    d.loc[nd, "flags"] += "NON_DATA_ROW;"
    tr = pyproj.Transformer.from_crs("EPSG:4326", "EPSG:32636", always_xy=True)
    e, n = tr.transform(d["lon"].fillna(0).values, d["lat"].fillna(0).values)
    d["utm_e"] = np.where(nd, np.nan, e)
    d["utm_n"] = np.where(nd, np.nan, n)
    d["utm_check_m"] = np.hypot(d["utm_e"] - d["x_file"], d["utm_n"] - d["y_file"])
    d.loc[~nd & d["x_file"].isna(), "flags"] += "UTM_FROM_LATLON;"
    d["diurnal_corr_nT"] = d["tmi_corr_nT"] - d["reading_nT"]
    # day segments: time steps backwards by more than 30 min start a new day
    t = d["time_hhmmss"].fillna(-1).astype(int)
    sec = np.where(nd, np.nan, (t // 10000) * 3600 + (t // 100 % 100) * 60 + t % 100)
    d["time_s"] = sec
    day = np.zeros(len(d), int)
    k, last = 1, None
    for i, s in enumerate(sec):
        if np.isnan(s):
            day[i] = 0
            continue
        if last is not None and s < last - 1800:
            k += 1
        day[i] = k
        last = s
    d["day_seq"] = day
    return d


def despike(d):
    """Hampel filter on the corrected total field, applied within each day."""
    d = d.copy()
    d["hampel_median_nT"] = np.nan
    d["hampel_dev_nT"] = np.nan
    for k, g in d[d["day_seq"] > 0].groupby("day_seq"):
        s = g["tmi_corr_nT"]
        med = s.rolling(HAMPEL_WINDOW, center=True, min_periods=3).median()
        mad = (s - med).abs().rolling(HAMPEL_WINDOW, center=True, min_periods=3).median()
        thr = np.maximum(HAMPEL_FLOOR_NT, HAMPEL_K * 1.4826 * mad)
        dev = (s - med).abs()
        d.loc[g.index, "hampel_median_nT"] = med
        d.loc[g.index, "hampel_dev_nT"] = dev
        spike = dev > thr
        d.loc[g.index[spike.values], "flags"] += "SPIKE_HAMPEL;"
    d["accepted"] = ~d["flags"].str.contains("NON_DATA_ROW|SPIKE_HAMPEL")
    return d


def add_igrf(d, date=SURVEY_DATE):
    import ppigrf

    d = d.copy()
    ok = d["day_seq"] > 0
    h_km = d.loc[ok, "elev_file"].fillna(d["elev_file"].median()).values / 1000.0
    be, bn, bu = ppigrf.igrf(d.loc[ok, "lon"].values, d.loc[ok, "lat"].values, h_km, date)
    be, bn, bu = (np.asarray(v).ravel() for v in (be, bn, bu))
    F = np.sqrt(be ** 2 + bn ** 2 + bu ** 2)
    d.loc[ok, "igrf_F_nT"] = F
    d.loc[ok, "igrf_inc_deg"] = np.degrees(np.arctan2(-bu, np.hypot(be, bn)))
    d.loc[ok, "igrf_dec_deg"] = np.degrees(np.arctan2(be, bn))
    d["tma_nT"] = d["tmi_corr_nT"] - d["igrf_F_nT"]
    return d


def igrf_date_sensitivity(lon, lat):
    import ppigrf

    out = []
    for date in (dt.datetime(2025, 1, 15), dt.datetime(2025, 7, 15), SURVEY_DATE, dt.datetime(2026, 7, 15)):
        be, bn, bu = ppigrf.igrf(lon, lat, 0.3, date)
        F = float(np.sqrt(be ** 2 + bn ** 2 + bu ** 2).ravel()[0])
        out.append(dict(date=date.date().isoformat(), F_nT=F,
                        inc_deg=float(np.degrees(np.arctan2(-bu, np.hypot(be, bn))).ravel()[0]),
                        dec_deg=float(np.degrees(np.arctan2(be, bn)).ravel()[0])))
    return pd.DataFrame(out)


def crossovers(d, maxdist=CROSSOVER_M):
    """Accepted readings < maxdist apart but measured on different days."""
    from scipy.spatial import cKDTree

    a = d[d["accepted"]].reset_index(drop=True)
    tree = cKDTree(a[["utm_e", "utm_n"]].values)
    pairs = tree.query_pairs(maxdist, output_type="ndarray")
    p = pairs[a["day_seq"].values[pairs[:, 0]] != a["day_seq"].values[pairs[:, 1]]]
    rows = pd.DataFrame({"day_a": a["day_seq"].values[p[:, 0]], "day_b": a["day_seq"].values[p[:, 1]],
                         "dTMI_nT": a["tmi_corr_nT"].values[p[:, 1]] - a["tmi_corr_nT"].values[p[:, 0]],
                         "dist_m": np.hypot(*(a[["utm_e", "utm_n"]].values[p[:, 1]] - a[["utm_e", "utm_n"]].values[p[:, 0]]).T)})
    swap = rows["day_a"] > rows["day_b"]
    rows.loc[swap, ["day_a", "day_b"]] = rows.loc[swap, ["day_b", "day_a"]].values
    rows.loc[swap, "dTMI_nT"] *= -1
    summ = rows.groupby(["day_a", "day_b"]).agg(n_pairs=("dTMI_nT", "size"), median_dTMI_nT=("dTMI_nT", "median"),
                                                mad_dTMI_nT=("dTMI_nT", lambda v: (v - v.median()).abs().median())
                                                ).reset_index()
    return rows, summ


def day_summary(d):
    g = d[d["day_seq"] > 0].groupby("day_seq")
    return g.agg(n_readings=("tmi_corr_nT", "size"), n_accepted=("accepted", "sum"),
                 time_first=("time_hhmmss", "first"), time_last=("time_hhmmss", "last"),
                 lon_median=("lon", "median"), lat_median=("lat", "median"),
                 diurnal_min_nT=("diurnal_corr_nT", "min"), diurnal_max_nT=("diurnal_corr_nT", "max"),
                 tmi_median_nT=("tmi_corr_nT", "median"), tmi_min_nT=("tmi_corr_nT", "min"),
                 tmi_max_nT=("tmi_corr_nT", "max")).reset_index()


# ----------------------------------------------------------------------------- gridding / RTP
def _chain():
    import verde as vd
    return vd.Chain([("block", vd.BlockReduce(np.median, spacing=BLOCK_M)),
                     ("trend", vd.Trend(degree=1)),
                     ("spline", vd.Spline(damping=1e-6, mindist=25.0))])


def _region(extent, pad_m=0.0):
    import pyproj
    tr = pyproj.Transformer.from_crs("EPSG:4326", "EPSG:32636", always_xy=True)
    lon0, lon1, lat0, lat1 = extent
    xs, ys = tr.transform([lon0, lon1, lon0, lon1], [lat0, lat0, lat1, lat1])
    return (min(xs) - pad_m, max(xs) + pad_m, min(ys) - pad_m, max(ys) + pad_m)


def _lonlat(g):
    import pyproj
    inv = pyproj.Transformer.from_crs("EPSG:32636", "EPSG:4326", always_xy=True)
    E, N = np.meshgrid(g.easting.values, g.northing.values)
    return inv.transform(E, N)


def _dist_to_data(g, x, y):
    from scipy.spatial import cKDTree
    E, N = np.meshgrid(g.easting.values, g.northing.values)
    dist, _ = cKDTree(np.column_stack([x, y])).query(np.column_stack([E.ravel(), N.ravel()]))
    return dist.reshape(E.shape)


def grid_field(x, y, v, extent=maps.EXTENT):
    chain = _chain().fit((x, y), v)
    g = chain.grid(region=_region(extent), spacing=GRID_SPACING_M, data_names="value", dims=("northing", "easting"))
    dist = _dist_to_data(g, x, y)
    g["value"] = g["value"].where(dist <= MASK_KM * 1000)
    return g


def cv_rmse(x, y, v):
    import verde as vd
    bx, bv = vd.BlockReduce(np.median, spacing=BLOCK_M).filter((x, y), v)
    cv = vd.BlockKFold(spacing=2000, n_splits=5, shuffle=True, random_state=0)
    return float(-np.mean(vd.cross_val_score(_chain(), bx, bv, cv=cv, scoring="neg_root_mean_squared_error")))


def rtp_grid(x, y, tma, inc, dec, extent=maps.EXTENT):
    """Reduction to the pole of the total-field anomaly (induced magnetisation)."""
    import harmonica as hm
    import xrft

    pad_m = (MASK_KM + TAPER_KM) * 1000
    chain = _chain().fit((x, y), tma)
    g = chain.grid(region=_region(extent, pad_m), spacing=GRID_SPACING_M, data_names="value",
                   dims=("northing", "easting"))["value"]
    dist = _dist_to_data(g.to_dataset(), x, y)
    level = float(np.median(tma))
    w = np.clip((MASK_KM * 1000 + TAPER_KM * 1000 - dist) / (TAPER_KM * 1000), 0, 1)
    w = np.sin(0.5 * np.pi * w) ** 2                         # 1 inside the data mask, 0 beyond the taper
    dev = (g - level) * w
    n_pad = int(np.ceil(0.5 * max(dev.shape)))
    padded = xrft.pad(dev, {"easting": n_pad, "northing": n_pad})
    r = hm.reduction_to_pole(padded, inclination=inc, declination=dec)
    r = xrft.unpad(r, {"easting": n_pad, "northing": n_pad})
    r = r.assign_coords(easting=g.easting, northing=g.northing)
    r = r.where(dist <= MASK_KM * 1000)
    reg = _region(extent)
    r = r.sel(easting=slice(reg[0], reg[1]), northing=slice(reg[2], reg[3]))
    return r.to_dataset(name="value")


# ----------------------------------------------------------------------------- figures
def equalised_norm(values, n_levels=64):
    """Histogram-equalised colour classes (equal numbers of values per colour band)."""
    v = np.asarray(values, float)
    v = v[np.isfinite(v)]
    levels = np.unique(np.round(np.percentile(v, np.linspace(0.5, 99.5, n_levels + 1)), 1))
    return BoundaryNorm(levels, ncolors=256, extend="both"), levels


def _equalised_colorbar(fig, m, ax, label, levels):
    cb = maps.colorbar(fig, m, ax, label + "  (histogram-equalised)")
    idx = np.linspace(0, len(levels) - 1, 9).round().astype(int)
    cb.set_ticks(levels[idx])
    cb.set_ticklabels([f"{levels[i]:,.0f}" for i in idx])
    return cb


def _credit(ax, text):
    maps.credit(ax, "\n".join(textwrap.wrap(text, 175)))


def _nice_step(lo, hi, n=18):
    raw = (hi - lo) / n
    mag = 10 ** np.floor(np.log10(raw))
    for m in (1, 2, 2.5, 5, 10):
        if raw <= m * mag:
            return m * mag
    return 10 * mag


def map_points(bm, d, ref, col, title, label, fname, fdir, center=None, subtitle="", note=""):
    a = d[d["accepted"]]
    norm, levels = equalised_norm(a[col])
    fig, ax = plt.subplots(figsize=(8.6, 5.6))
    bm.draw(ax, relief_alpha=0.85)
    o = a.iloc[np.argsort(np.abs(a[col].values - np.median(a[col])))]
    m = ax.scatter(o["lon"], o["lat"], c=o[col], s=3.2, cmap=maps.ANOM, norm=norm, lw=0, zorder=12, rasterized=True)
    sp = d[d["flags"].str.contains("SPIKE")]
    ax.scatter(sp["lon"], sp["lat"], s=7, marker="x", c="#1d1d1b", lw=0.5, zorder=13,
               label=f"rejected spikes ({len(sp)})")
    maps.base100(ax, ref)
    maps.scalebar(ax)
    maps.north_arrow(ax)
    _equalised_colorbar(fig, m, ax, label, levels)
    ax.legend(loc="lower right", markerscale=1.5)
    ax.set_title(title, loc="left", pad=14 if subtitle else 6)
    if subtitle:
        ax.text(0, 1.012, subtitle, transform=ax.transAxes, fontsize=7.2, color=maps.INK2, va="bottom")
    _credit(ax, note)
    fig.savefig(os.path.join(fdir, fname), dpi=maps.DPI)
    plt.close(fig)


def map_grid(bm, g, d, ref, title, label, fname, fdir, center=None, subtitle="", note="", track_alpha=0.45):
    LON, LAT = _lonlat(g)
    Z = g["value"].values
    lo, hi = np.nanpercentile(Z, [1, 99])
    norm, levels = equalised_norm(Z)
    step = _nice_step(lo, hi)
    fig, ax = plt.subplots(figsize=(8.6, 5.6))
    bm.draw(ax, relief_alpha=0.85)
    m = ax.pcolormesh(LON, LAT, Z, cmap=maps.ANOM, norm=norm, shading="auto", alpha=0.85, zorder=5, rasterized=True)
    lev = np.arange(np.floor(lo / step) * step, np.ceil(hi / step) * step + step / 2, step)
    cs = ax.contour(LON, LAT, Z, levels=lev, colors=maps.INK, linewidths=0.35, alpha=0.7, zorder=6)
    lab = ax.clabel(cs, levels=cs.levels[::2], fmt="%g", fontsize=5.5, inline=True, inline_spacing=2)
    for t in lab:
        t.set_path_effects([pe.withStroke(linewidth=1.6, foreground="white")])
    a = d[d["accepted"]]
    ax.scatter(a["lon"].values[::4], a["lat"].values[::4], s=0.25, c=maps.INK, lw=0, alpha=track_alpha, zorder=7,
               rasterized=True)
    maps.base100(ax, ref)
    maps.scalebar(ax)
    maps.north_arrow(ax)
    _equalised_colorbar(fig, m, ax, label, levels)
    ax.set_title(title, loc="left", pad=14 if subtitle else 6)
    if subtitle:
        ax.text(0, 1.012, subtitle, transform=ax.transAxes, fontsize=7.2, color=maps.INK2, va="bottom")
    _credit(ax, f"Contour interval {step:g} nT. Biharmonic spline after {BLOCK_M:.0f} m block median, "
                    f"{GRID_SPACING_M:.0f} m grid, blanked >{MASK_KM:g} km from data; grey dots = survey tracks. "
                    f"Relief: Copernicus GLO-30. {note}")
    fig.savefig(os.path.join(fdir, fname), dpi=maps.DPI)
    plt.close(fig)


def save_grid(g, path_stem, name):
    LON, LAT = _lonlat(g)
    g.to_netcdf(path_stem + ".nc")
    E, N = np.meshgrid(g.easting, g.northing)
    pd.DataFrame({"lon": LON.ravel(), "lat": LAT.ravel(), "easting_utm36n": E.ravel(), "northing_utm36n": N.ravel(),
                  name: g["value"].values.ravel()}).dropna().to_csv(path_stem + "_xyz.csv", index=False,
                                                                     float_format="%.6f")


def write_workbook(path, d, days, xs, igrf_sens, meta):
    from .export import _write_sheet
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [("Wadi Ghadir ground magnetic survey – processed total field, anomaly and QC", True),
             ("Prepared by Dr. Mohamed Sobh, LIAG. Produced by scripts/run_magnetics.py.", False), ("", False)]
    lines += [(f"{k}: {v}", False) for k, v in meta.items()]
    lines += [("", False), ("Columns of Mag_Data", True),
              ("tmi_corr_nT = diurnally corrected total field as supplied; igrf_F_nT = IGRF-14 total field; "
               "tma_nT = tmi_corr_nT - igrf_F_nT (total-field anomaly); accepted = not a spike/non-data row; "
               "flags = QC reasons; day_seq = survey day inferred from the time stamps.", False)]
    for t, b in lines:
        ws.append([t])
        if b:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
    ws.column_dimensions["A"].width = 150
    cols = ["source_row", "day_seq", "time_hhmmss", "lat", "lon", "utm_e", "utm_n", "elev_file", "reading_nT",
            "diurnal_corr_nT", "tmi_corr_nT", "igrf_F_nT", "igrf_inc_deg", "igrf_dec_deg", "tma_nT",
            "hampel_dev_nT", "accepted", "flags"]
    _write_sheet(wb.create_sheet("Mag_Data"), d[cols], {"lat": "0.0000000", "lon": "0.0000000", "utm": "0.00",
                                                       "_nT": "0.00", "_deg": "0.000"})
    _write_sheet(wb.create_sheet("Day_Summary"), days)
    _write_sheet(wb.create_sheet("Crossovers"), xs)
    _write_sheet(wb.create_sheet("IGRF_Date_Sensitivity"), igrf_sens)
    wb.save(path)
