"""Deliverable exports: formatted Excel workbook, GeoJSON and KML of the final stations."""
import json
import os

import numpy as np
import pandas as pd

from .pipeline import RHO_LIST

# (output column, source column, unit / description)
FINAL_COLUMNS = [
    ("Station_UID", "occ_key", "instrument|date|line|station of the occupation"),
    ("Instrument", "instrument", "CG-6 serial label"),
    ("Date", "date", "survey date (UTC)"),
    ("Line", "line", "instrument line number"),
    ("Station", "station", "instrument station number (not unique across lines/instruments)"),
    ("GNSS_Point", "gnss_point_id", "matched GNSS point label"),
    ("GNSS_Source", "gnss_source", "file the GNSS point came from"),
    ("Match_Dist_m", "match_dist_m", "CG-6 GPS to GNSS point separation [m]"),
    ("Lat_deg", "lat", "latitude, WGS84 [deg]"),
    ("Lon_deg", "lon", "longitude, WGS84 [deg]"),
    ("UTM36N_E_m", "utm_e", "UTM zone 36N easting [m]"),
    ("UTM36N_N_m", "utm_n", "UTM zone 36N northing [m]"),
    ("Time_UTC", "t_mid", "mean time of the accepted readings (UTC)"),
    ("N_Readings", "n_accepted", "accepted 60-s readings"),
    ("h_ellip_m", "h_used", "GNSS ellipsoidal height incl. day-7 adjustment [m]"),
    ("N_EGM2008_m", "N_egm2008_m", "EGM2008 geoid undulation [m]"),
    ("H_ortho_m", "H_ortho_m", "orthometric height H = h - N [m]"),
    ("H_DEM_m", "H_dem_m", "Copernicus GLO-30 height at the station [m]"),
    ("H_minus_DEM_m", "H_minus_dem_m", "H_ortho - H_DEM [m] (QC)"),
    ("dH_rel_base100_m", "dh_m", "H - H(base 100) [m]"),
    ("dg_rel_base100_mGal", "dg100_scaled_mgal", "observed gravity minus base 100 [mGal] (#0313 x k)"),
    ("sigma_dg_mGal", "dg100_sigma_mgal", "1-sigma of dg [mGal]"),
    ("dNormalGravity_mGal", "normal_gravity_diff_mgal", "gamma(phi) - gamma(phi100), WGS84 [mGal]"),
    ("dFreeAirTerm_mGal", "free_air_term_diff_mgal", "FA(H) - FA(H100), 2nd order [mGal]"),
    ("FA_rel_mGal", "dFA_rel_mgal", "free-air anomaly relative to base 100 [mGal]"),
    ("sigma_FA_mGal", "dFA_sigma_mgal", "1-sigma of FA [mGal]"),
    ("g1_topo_mGal_per_gcc", "g1_topo_mgal_per_gcc", "DEM topographic attraction for 1 g/cm3 [mGal]"),
    ("u_topo_rel_base100", "u_topo_mgal_per_gcc", "g1(station) - g1(base 100) [mGal per g/cm3]"),
    ("TC_2.67_mGal", "TC_rho2.67_mgal", "terrain + curvature correction rel. to slab, rho 2.67 [mGal]"),
]
for _r in RHO_LIST:
    FINAL_COLUMNS.append((f"SB_rel_{_r:.2f}_mGal", f"dSB_rel_rho{_r:.2f}_mgal",
                          f"simple Bouguer anomaly rel. base 100, rho {_r:.2f} g/cm3 [mGal]"))
for _r in RHO_LIST:
    FINAL_COLUMNS.append((f"CB_rel_{_r:.2f}_mGal", f"dCB_rel_rho{_r:.2f}_mgal",
                          f"complete Bouguer anomaly rel. base 100, rho {_r:.2f} g/cm3 [mGal]"))
FINAL_COLUMNS += [
    ("Dist_to_Sea_km", "dist_to_sea_km", "distance to Red Sea coast (DEM) [km]; bathymetry not modelled"),
    ("Drift_Mode", "drift_mode", "how the base-100 reference was obtained"),
    ("Scale_k_applied", "scale_factor_applied", "instrument scale factor applied"),
    ("Day_Height_Shift_m", "day_height_shift_m", "GNSS day datum adjustment subtracted [m]"),
    ("Match_Status", "match_status", "GNSS match outcome"),
    ("QC_Warnings", "warnings", "row-level warnings (not rejections)"),
]


def final_station_table(st):
    import pyproj
    s = st[st["product_ok"]].copy()
    tr = pyproj.Transformer.from_crs("EPSG:4326", "EPSG:32636", always_xy=True)
    s["utm_e"], s["utm_n"] = tr.transform(s["lon"].values, s["lat"].values)
    s["t_mid"] = pd.to_datetime(s["t_mid"]).dt.strftime("%Y-%m-%d %H:%M:%S")
    out = pd.DataFrame({o: s[c].values for o, c, _ in FINAL_COLUMNS})
    return out.sort_values(["Instrument", "Date", "Time_UTC"]).reset_index(drop=True)


def _write_sheet(ws, df, number_formats=None, widths=None):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1C3D5A")
    head_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="D0D0D0")
    ws.append(list(df.columns))
    for c in ws[1]:
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and not np.isfinite(v)) else
                   (v.item() if isinstance(v, np.generic) else v) for v in row])
    ws.freeze_panes = "B2" if df.shape[1] > 6 else "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for j, col in enumerate(df.columns, start=1):
        letter = get_column_letter(j)
        w = (widths or {}).get(col) or min(max(len(str(col)) * 0.9, 9), 28)
        ws.column_dimensions[letter].width = w
        fmt = None
        for key, f in (number_formats or {}).items():
            if key in col:
                fmt = f
        if fmt is None and pd.api.types.is_float_dtype(df[col]):
            fmt = "0.000"
        if fmt:
            for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2, max_row=ws.max_row):
                for c in cell:
                    c.number_format = fmt


def write_workbook(path, ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    final = final_station_table(ctx["st"])
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Wadi Ghadir gravity survey (Eastern Desert, Egypt), 16-24 January 2026 – final results", True),
        ("Prepared by Dr. Mohamed Sobh, LIAG. Produced by scripts/run_processing.py from the raw CG-6 and GNSS files.", False),
        ("", False),
        ("Conventions", True),
        ("All anomalies are RELATIVE TO BASE 100 (24.786749 N, 34.933755 E). No absolute gravity tie exists; add the "
         "absolute anomaly of base 100 when it becomes available.", False),
        ("Heights: GNSS ellipsoidal heights (shown by DEM comparison) converted to orthometric heights with EGM2008.", False),
        ("Free-air: second-order normal-gravity gradient; normal gravity: WGS84 Somigliana.", False),
        ("Simple Bouguer: infinite slab 0.04193*rho*dH. Complete Bouguer: Copernicus GLO-30 prisms to 22 km incl. Earth "
         "curvature (Harmonica); Red Sea water/bathymetry not modelled.", False),
        ("CG-6 #0313 values are multiplied by the inter-instrument scale factor k (sheet Instrument_Scale).", False),
        ("Recommended map product: CB_rel_2.67_mGal (provisional density; see Density_Consensus).", False),
        ("", False),
        ("Sheets", True),
        ("Final_Stations – one row per accepted field station: coordinates, heights, gravity, FA, SB and CB for 11 densities", False),
        ("Column_Dictionary – definition and unit of every Final_Stations column", False),
        ("Profiles – stations along each Nettleton traverse with chainage (for profile plots)", False),
        ("Nettleton_Results / Nettleton_Sweeps / Nettleton_Windows / Density_Consensus – density analysis", False),
        ("Base_Control / Hotel_Base_Ties / Daily_Closures / Instrument_Scale – base control and instrument reconciliation", False),
        ("Height_Datum_Check / Terrain_Validation – geodetic and terrain-model checks", False),
        ("Occupations / Rows_QC – every occupation and every one of the 1519 raw readings with QC decision", False),
        ("Instrument_Metadata – CG-6 header information", False),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=12 if ws.max_row == 1 else 11)
    ws.column_dimensions["A"].width = 150

    fmts = {"_deg": "0.000000", "UTM": "0", "_m": "0.00", "mGal": "0.000", "sigma": "0.000", "km": "0.0"}
    _write_sheet(wb.create_sheet("Final_Stations"), final, fmts, widths={"Station_UID": 30, "GNSS_Source": 22})
    _write_sheet(wb.create_sheet("Column_Dictionary"),
                 pd.DataFrame([(o, d) for o, _, d in FINAL_COLUMNS], columns=["Column", "Definition / unit"]),
                 widths={"Column": 24, "Definition / unit": 90})
    pcols = ["profile_id", "instrument", "date", "line", "station", "chainage_m", "lat", "lon", "H_ortho_m", "H_dem_m",
             "dh_m", "dFA_rel_mgal", "u_slab_mgal_per_gcc", "u_topo_mgal_per_gcc", "dSB_rel_rho2.67_mgal",
             "dCB_rel_rho2.67_mgal", "drift_mode"]
    _write_sheet(wb.create_sheet("Profiles"), ctx["prof"][pcols].sort_values(["profile_id", "chainage_m"]),
                 {"lat": "0.000000", "lon": "0.000000"})
    _write_sheet(wb.create_sheet("Nettleton_Results"), ctx["net"])
    _write_sheet(wb.create_sheet("Nettleton_Sweeps"), ctx["sweeps"][["profile_id", "method", "rho", "corr_detrended",
                                                                     "rms_detrended_mgal", "roughness_2nd_diff_mgal"]])
    _write_sheet(wb.create_sheet("Nettleton_Windows"), ctx["win"])
    _write_sheet(wb.create_sheet("Density_Consensus"), ctx["cons"])
    _write_sheet(wb.create_sheet("Base_Control"), ctx["ctrl"].assign(t=lambda d: d["t"].astype(str)))
    _write_sheet(wb.create_sheet("Hotel_Base_Ties"), ctx["ties"].assign(t_hotel=lambda d: d["t_hotel"].astype(str),
                                                                        t_base100=lambda d: d["t_base100"].astype(str)))
    _write_sheet(wb.create_sheet("Daily_Closures"), ctx["closures"].assign(
        base100_first_time=lambda d: d["base100_first_time"].astype(str)))
    sc = {k: v for k, v in ctx["sc"].items() if not isinstance(v, dict)}
    _write_sheet(wb.create_sheet("Instrument_Scale"), pd.concat(
        [pd.DataFrame([sc]).T.reset_index().rename(columns={"index": "quantity", 0: "value"}),
         pd.DataFrame([{}]), ctx["pairs"]], ignore_index=True))
    _write_sheet(wb.create_sheet("Height_Datum_Check"), ctx["hchk"])
    _write_sheet(wb.create_sheet("Terrain_Validation"), ctx["tval"])
    occ = ctx["occ"].copy()
    for c in ("t_start", "t_end", "t_mid"):
        occ[c] = occ[c].astype(str)
    _write_sheet(wb.create_sheet("Occupations"), occ)
    rows = ctx["df"].drop(columns=["t"], errors="ignore").copy()
    rows["time_utc"] = rows["time_utc"].astype(str)
    _write_sheet(wb.create_sheet("Rows_QC"), rows)
    _write_sheet(wb.create_sheet("Instrument_Metadata"), ctx["meta"].T.reset_index().rename(
        columns={"index": "field"}).astype(str))
    wb.save(path)
    return final


def write_geojson(path, final):
    feats = []
    keep = ["Station_UID", "Instrument", "Date", "Line", "Station", "H_ortho_m", "dg_rel_base100_mGal",
            "FA_rel_mGal", "SB_rel_2.67_mGal", "CB_rel_2.67_mGal", "TC_2.67_mGal", "sigma_FA_mGal"]
    for _, r in final.iterrows():
        props = {k: (None if pd.isna(r[k]) else (round(float(r[k]), 4) if isinstance(r[k], (float, np.floating)) else r[k]))
                 for k in keep}
        feats.append({"type": "Feature", "geometry": {"type": "Point", "coordinates": [round(r["Lon_deg"], 7),
                                                                                       round(r["Lat_deg"], 7)]},
                      "properties": props})
    with open(path, "w") as fh:
        json.dump({"type": "FeatureCollection", "name": "wadi_ghadir_gravity_stations",
                   "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:OGC:1.3:CRS84"}},
                   "features": feats}, fh)


def write_kml(path, final):
    col = {"CG6-0640": "ffd6782a", "CG6-0313": "ff3468eb"}
    out = ['<?xml version="1.0" encoding="UTF-8"?>', '<kml xmlns="http://www.opengis.net/kml/2.2"><Document>',
           "<name>Wadi Ghadir gravity stations</name>"]
    for inst, c in col.items():
        out.append(f'<Style id="{inst}"><IconStyle><color>{c}</color><scale>0.6</scale><Icon><href>'
                   "http://maps.google.com/mapfiles/kml/shapes/placemark_circle.png</href></Icon></IconStyle>"
                   "<LabelStyle><scale>0</scale></LabelStyle></Style>")
    for inst, d in final.groupby("Instrument"):
        out.append(f"<Folder><name>{inst}</name>")
        for _, r in d.iterrows():
            desc = (f"Date {r['Date']}, line {r['Line']}, station {r['Station']}<br/>H = {r['H_ortho_m']:.1f} m<br/>"
                    f"FA = {r['FA_rel_mGal']:.2f} mGal<br/>CB(2.67) = {r['CB_rel_2.67_mGal']:.2f} mGal (rel. base 100)")
            out.append(f"<Placemark><name>{inst[-4:]} L{r['Line']}-{r['Station']}</name><styleUrl>#{inst}</styleUrl>"
                       f"<description><![CDATA[{desc}]]></description><Point><coordinates>{r['Lon_deg']:.7f},"
                       f"{r['Lat_deg']:.7f},0</coordinates></Point></Placemark>")
        out.append("</Folder>")
    out.append("</Document></kml>")
    with open(path, "w") as fh:
        fh.write("\n".join(out))
