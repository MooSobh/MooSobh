"""Formatted Excel sheets (openpyxl): bold frozen header, number formats, column widths.

Author : Dr. Mohamed Sobh, LIAG (Leibniz Institute for Applied Geophysics)
Project: Wadi Ghadir gravity and magnetic survey, Eastern Desert, Egypt (2026)
"""
import numpy as np
import pandas as pd


def write_sheet(ws, df, number_formats=None, widths=None):
    """Write a DataFrame to an openpyxl worksheet: bold frozen header, number formats by column-name pattern, column widths."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1C3D5A")
    head_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="D0D0D0")
    ws.append(list(df.columns))
    for c in ws[1]:
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and not np.isfinite(v)) else
                   (v.item() if isinstance(v, np.generic) else v) for v in row])
    ws.freeze_panes = "B2" if df.shape[1] > 6 else "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for j, col in enumerate(df.columns, start=1):
        letter = get_column_letter(j)
        w = (widths or {}).get(col) or min(max(len(str(col)) * 0.9, 9), 28)
        ws.column_dimensions[letter].width = w
        fmt = None
        for key, f in (number_formats or {}).items():
            if key in col:
                fmt = f
        if fmt is None and pd.api.types.is_float_dtype(df[col]):
            fmt = "0.000"
        if fmt:
            for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2, max_row=ws.max_row):
                for c in cell:
                    c.number_format = fmt
